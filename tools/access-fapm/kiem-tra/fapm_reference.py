"""Ban tinh doi chieu FA cho PM/PD (Muc A.I - ALL/AD/2210-006).

Cung logic voi modFAPM.bas. Chay: python3 fapm_reference.py
  1. Kiem tra cong thuc voi cac kich ban mau.
  2. Sinh FAPM_Input_Mau.xlsx (sheet Agent 98 cot, Holiday, FAPM_Input, KetQua_KyVong).
"""

import calendar
import json
import sys
from datetime import date, timedelta
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "access-fa" / "kiem-tra"))
from fa_reference import AGENT_HEADERS, DATE_HEADERS  # noqa: E402  (mau Agent 98 cot dung chung)

# ---- Tham so (khop tblParam trong Build_FAPM_Tool.vbs) ----
STD_DAYS = 22
MIN_ATT_PCT = 0          # 0 = khong ap dung dieu kien cham cong toi thieu (van ban 006 khong neu)
MIN_DAYS = 0
K2_NA_COEF = 1.0
NEW_GUAR_ALLOW = 500_000     # PM <= 12 thang ke tu thang tien: >= 8 luot LP hoat dong trong quy xet
NEW_GUAR_MONTHS = 12
NEW_GUAR_MIN_TURNS = 8
QTR_GUAR_ALLOW = 150_000     # PM khac: >= 6 luot LP hoat dong trong quy xet
QTR_GUAR_MIN_TURNS = 6
JOIN_GUAR_ALLOW = 500_000    # PM gia nhap den het Q2/2027: 6 thang dau
JOIN_GUAR_MONTHS = 6
JOIN_GUAR_DEADLINE = "202706"
FA_TITLES = {"PM", "PD"}

# (FYC nhom truc tiep toi thieu, so LP hoat dong toi thieu, phu cap/ngay)
ALLOWANCE = [(10_000_000, 2, 300_000), (30_000_000, 2, 500_000), (60_000_000, 2, 800_000)]
K2_TIERS = [(80, 1.2), (70, 1.1), (60, 1.0), (0, 0.5)]
LPCOUNT_TIERS = [(100, 1.15), (50, 1.05), (0, 1.0)]


def rhu(x, d=0):
    return float(Decimal(str(x)).quantize(Decimal(1).scaleb(-d), rounding=ROUND_HALF_UP))


def work_days(d1, d2, holidays):
    n, d = 0, d1
    while d <= d2:
        if d.weekday() < 5 and d not in holidays:
            n += 1
        d += timedelta(days=1)
    return n


def allowance(group_fyc, active_lp):
    return max([a for f, n, a in ALLOWANCE if group_fyc >= f and active_lp >= n] + [0])


def k2_coef(k2):
    if k2 is None:
        return K2_NA_COEF
    k = rhu(k2, 1)
    return next(c for m, c in K2_TIERS if k >= m)


def lpcount_coef(turns):
    return next(c for m, c in LPCOUNT_TIERS if turns >= m)


def months_between(p1, p2):
    """So thang tu YYYYMM p1 den p2 (p1=p2 -> 1)."""
    return (int(p2[:4]) - int(p1[:4])) * 12 + int(p2[4:]) - int(p1[4:]) + 1


def calc(period, agent, inp, holidays, prior_turns=None, qtr_turns=0, qtr_prev=None, ky=1):
    """Tinh FA PM/PD 1 thang.
    prior_turns: {period: ActiveLP} cac thang truoc (de cong don luot LP ke tu thang tien)
    qtr_turns: tong luot LP hoat dong trong quy xet (quy truoc)
    qtr_prev: [(FA_Month, QtrGuarMonth)] cua 2 thang dau quy nay (chi dung khi thang 3 va ky 4)
    """
    y, m = int(period[:4]), int(period[4:])
    start, end = date(y, m, 1), date(y, m, calendar.monthrange(y, m)[1])
    appoint, promo = agent["AppointDate"], agent.get("PromoDate") or agent["AppointDate"]
    wm = (y - appoint.year) * 12 + (m - appoint.month) + 1
    promo_period = promo.strftime("%Y%m")
    pm_months = months_between(promo_period, period)

    base = work_days(appoint if wm == 1 else start, end, holidays)
    att = min(1.0, (inp["ActualDays"] + 1) / base) if base else 0
    att_pct = rhu(att * 100, 1)
    day_f = rhu(min(inp["ActualDays"], STD_DAYS), 1) if wm == 1 else rhu(att_pct / 100 * STD_DAYS, 1)

    reason = ""
    title = (agent.get("Title") or "").upper()
    if title and title not in FA_TITLES:
        reason = "Chức danh không phải PM/PD"
    elif agent.get("TerminateDate") and agent["TerminateDate"] <= end:
        reason = "HĐ đại lý chấm dứt trong/trước tháng tính"
    elif (inp.get("StatusEOM") or "").upper() == "SE":
        reason = "Tình trạng SE cuối tháng"
    elif att_pct < MIN_ATT_PCT and inp["ActualDays"] < MIN_DAYS:
        reason = "Không đạt chấm công tối thiểu"
    eligible = reason == ""

    cum_turns = inp["ActiveLP"] + sum(v for p, v in (prior_turns or {}).items() if promo_period <= p < period)
    allow = allowance(inp["GroupFYC"], inp["ActiveLP"])
    k2c, lpc = k2_coef(inp["K2Pct"]), lpcount_coef(cum_turns)
    fa_formula = rhu(allow * day_f * k2c * lpc)

    join_guar = 0
    if eligible and promo_period <= JOIN_GUAR_DEADLINE and 1 <= pm_months <= JOIN_GUAR_MONTHS:
        join_guar = rhu(JOIN_GUAR_ALLOW * day_f)
    fa_month = max(fa_formula, join_guar) if eligible else 0

    # Bao dam quy: xet o cuoi quy truoc theo luot LP hoat dong va tham nien PM tai thoi diem do
    month_in_qtr = (m - 1) % 3 + 1
    prev_qtr_end = f"{y}{m - month_in_qtr:02d}" if m > month_in_qtr else f"{y - 1}12"
    pm_months_at_review = months_between(promo_period, prev_qtr_end)
    level = 0
    if pm_months_at_review <= NEW_GUAR_MONTHS and qtr_turns >= NEW_GUAR_MIN_TURNS:
        level = NEW_GUAR_ALLOW
    elif qtr_turns >= QTR_GUAR_MIN_TURNS:
        level = QTR_GUAR_ALLOW
    qtr_month = rhu(level * day_f) if level and eligible else 0
    topup = 0
    if level and month_in_qtr == 3 and ky == 4:
        s_guar = qtr_month + sum(g for _, g in (qtr_prev or []))
        s_fa = fa_month + sum(f for f, _ in (qtr_prev or []))
        topup = max(0, s_guar - s_fa)

    return dict(WorkMonthNo=wm, PromoMonthNo=pm_months, WorkDaysBase=base, AttendancePct=att_pct, DayFactor=day_f,
                Eligible=eligible, Reason=reason, GroupFYC=inp["GroupFYC"], ActiveLP=inp["ActiveLP"],
                CumLPTurns=cum_turns, AllowancePerDay=allow, K2Coef=k2c, LPCountCoef=lpc, FA_Formula=fa_formula,
                JoinGuar=join_guar, FA_Month=fa_month, QtrGuarLevel=level, QtrGuarMonth=qtr_month,
                QtrGuarTopup=topup, FA_Total=fa_month + topup)


# ---- Du lieu mau: thang 10/2026 ----
PERIOD = "202610"
HOLIDAYS = {date(2026, 9, 2): "Quốc khánh", date(2027, 1, 1): "Tết Dương lịch"}
AGENTS = [  # code, ten, ngay bo nhiem, ngay thang tien PM, chuc danh, kich ban
    ("PM001", "PM A", date(2026, 3, 2), None, "PM", "Mức 0,5tr, K2 N/A, lượt LP < 50 → 10.500.000"),
    ("PM002", "PM B", date(2025, 6, 2), date(2026, 4, 1), "PM", "0,8tr × 22 × K2 120% × lượt LP ≥ 50 (105%)"),
    ("PM003", "PM C", date(2026, 8, 3), None, "PM", "FYC nhóm < 10tr → công thức 0, bảo đảm gia nhập 0,5tr"),
    ("PM004", "PM D", date(2025, 9, 1), None, "PM", "SE cuối tháng → không chi"),
    ("PM005", "PM E", date(2025, 1, 6), None, "PM", "Chỉ 1 LP hoạt động → mức 0; quá 6 tháng gia nhập"),
    ("LP006", "LP F", date(2026, 1, 5), None, "LP", "Chức danh LP → không thuộc tool này"),
    ("PM007", "PM G", date(2025, 6, 2), None, "PM", "K2 55% → hệ số 50%"),
    ("PD008", "PD H", date(2024, 2, 1), None, "PD", "PD cũng được FA nhóm trực tiếp"),
]
INPUTS = {
    "PM001": dict(ActualDays=20, GroupFYC=35_000_000, ActiveLP=3, K2Pct=None, StatusEOM="Active"),
    "PM002": dict(ActualDays=22, GroupFYC=70_000_000, ActiveLP=4, K2Pct=85, StatusEOM="Active"),
    "PM003": dict(ActualDays=18, GroupFYC=8_000_000, ActiveLP=2, K2Pct=None, StatusEOM="Active"),
    "PM004": dict(ActualDays=18, GroupFYC=40_000_000, ActiveLP=3, K2Pct=72, StatusEOM="SE"),
    "PM005": dict(ActualDays=20, GroupFYC=40_000_000, ActiveLP=1, K2Pct=75, StatusEOM="Active"),
    "LP006": dict(ActualDays=20, GroupFYC=20_000_000, ActiveLP=2, K2Pct=None, StatusEOM="Active"),
    "PM007": dict(ActualDays=22, GroupFYC=12_000_000, ActiveLP=2, K2Pct=55, StatusEOM="Active"),
    "PD008": dict(ActualDays=21, GroupFYC=65_000_000, ActiveLP=5, K2Pct=90, StatusEOM="Active"),
}
# Luot LP hoat dong cac thang truoc (de test cong don): PM B 8 LP/thang tu 04/2026 -> 48 + 4 = 52
PRIOR = {"PM002": {f"2026{m:02d}": 8 for m in range(4, 10)}}


def main():
    hol = set(HOLIDAYS)
    res = {}
    for code, name, ap, promo, title, _ in AGENTS:
        agent = dict(AppointDate=ap, PromoDate=promo, Title=title)
        res[code] = calc(PERIOD, agent, INPUTS[code], hol, PRIOR.get(code))

    assert res["PM001"]["FA_Total"] == 10_500_000, res["PM001"]
    assert res["PM002"]["CumLPTurns"] == 52 and res["PM002"]["FA_Total"] == 22_176_000, res["PM002"]
    assert res["PM003"]["FA_Formula"] == 0 and res["PM003"]["FA_Total"] == 9_500_000, res["PM003"]
    assert res["PM004"]["FA_Total"] == 0 and res["PM005"]["FA_Total"] == 0 and res["LP006"]["FA_Total"] == 0
    assert res["PM007"]["FA_Total"] == 3_300_000, res["PM007"]
    assert res["PD008"]["FA_Total"] == 800_000 * 22 * 1.2, res["PD008"]

    # Bao dam quy (thang 3 cua quy, ky 4): PM moi, quy truoc 9 luot -> 0,5tr/ngay, bu phan thieu
    r = calc("202612", dict(AppointDate=date(2026, 3, 2), Title="PM"),
             dict(ActualDays=20, GroupFYC=5_000_000, ActiveLP=1, K2Pct=None, StatusEOM="Active"), hol,
             qtr_turns=9, qtr_prev=[(2_000_000, 10_500_000), (3_000_000, 11_000_000)], ky=4)
    assert r["QtrGuarLevel"] == 500_000 and r["QtrGuarMonth"] == rhu(500_000 * r["DayFactor"]), r
    assert r["QtrGuarTopup"] == (r["QtrGuarMonth"] + 10_500_000 + 11_000_000) - (r["FA_Month"] + 5_000_000), r
    # PM lau nam: 6 luot -> 0,15tr
    r2 = calc("202612", dict(AppointDate=date(2024, 3, 2), Title="PM"),
              dict(ActualDays=20, GroupFYC=0, ActiveLP=0, K2Pct=None, StatusEOM="Active"), hol, qtr_turns=6, ky=1)
    assert r2["QtrGuarLevel"] == 150_000 and r2["QtrGuarTopup"] == 0

    out = Path(__file__).resolve().parent.parent / "FAPM_Input_Mau.xlsx"
    build_workbook(out, res)
    for code, *_ in AGENTS:
        print(f"{code}: FA={res[code]['FA_Total']:>12,.0f}  {res[code]['Reason']}")
    print("OK ->", out)


def build_workbook(path, res):
    wb = Workbook()
    head, fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E78")

    def sheet(title, cols, rows, widths=None):
        ws = wb.create_sheet(title)
        ws.append(cols)
        for c in ws[1]:
            c.font, c.fill = head, fill
        for r in rows:
            ws.append(list(r))
        for i in range(len(cols)):
            ws.column_dimensions[get_column_letter(i + 1)].width = (widths or {}).get(i, 16)
        ws.freeze_panes = "A2"
        return ws

    hd = wb.active
    hd.title = "HuongDan"
    for line in [
        "MẪU NHẬP DỮ LIỆU TÍNH PHỤ CẤP CỐ ĐỊNH (FA) CHO PM/PD – ALL/AD/2210-006 mục A.I",
        "",
        "Tool đọc 3 sheet: Agent, Holiday, FAPM_Input. Không đổi tên sheet và tên cột.",
        "Sheet Agent: dán nguyên file xuất hệ thống (Tbl_Agent 98 cột). Tool dùng AGENT_CODE, FULL_NAME, TITLE (Code),",
        "  APPOINTMENT_DATE, PROMOTION_DATE (ngày thăng tiến PM; trống thì lấy APPOINTMENT_DATE), TERMINATION_DATE, AGENT_STATUS.",
        "Sheet FAPM_Input – số liệu nhóm trực tiếp từng tháng/kỳ (lũy kế đến ngày chốt ACK của kỳ):",
        "  Period (YYYYMM) | Ky (1-4) | AgentCode | GroupFYC: FYC nhóm trực tiếp phát sinh (VND)",
        "  ActiveLP: số LP nhóm trực tiếp hoạt động trong tháng (LP có ≥ 1 HĐ và ≥ 4 triệu FYC)",
        "  K2Pct: %K2 nhóm trực tiếp, ví dụ 85. Để TRỐNG nếu chưa có K2 (N/A)",
        "  StatusEOM: tình trạng cuối tháng (Active / SE). Để trống thì lấy AGENT_STATUS trong Agent",
        "Số ngày đi làm lấy từ Tbl_Attendance (Data chung). Lượt LP hoạt động lũy kế = tổng ActiveLP các tháng kể từ tháng thăng tiến.",
        "Sheet KetQua_KyVong: kết quả bản tính Python để đối soát, tool không đọc.",
    ]:
        hd.append([line])
    hd["A1"].font = Font(bold=True, size=13)
    hd.column_dimensions["A"].width = 130

    def agent_row(code, name, ap, promo, title):
        v = {"Period": PERIOD, "Channel": "AGENCY", "AGENT_CODE": code, "FULL_NAME": name, "TITLE (Code)": title,
             "AGENT_STATUS": "Active", "CHANNEL": "AGENCY", "AGENCY_TYPE": "FT", "APPOINTMENT_DATE": ap,
             "PROMOTION_DATE": promo}
        return [v.get(h) for h in AGENT_HEADERS]

    ws = sheet("Agent", AGENT_HEADERS, [agent_row(*a[:5]) for a in AGENTS])
    for i, h in enumerate(AGENT_HEADERS):
        fmt = "dd/mm/yyyy" if h in DATE_HEADERS else ("@" if h in ("Period", "AGENT_CODE") else None)
        if fmt:
            for row in ws.iter_rows(min_row=2, min_col=i + 1, max_col=i + 1):
                row[0].number_format = fmt

    ws = sheet("Holiday", ["HolidayDate", "Description"], HOLIDAYS.items(), {1: 30})
    for row in ws.iter_rows(min_row=2, max_col=1):
        row[0].number_format = "dd/mm/yyyy"

    cols = ["Period", "Ky", "AgentCode", "GroupFYC", "ActiveLP", "K2Pct", "StatusEOM"]
    rows = [[PERIOD, 1, c] + [INPUTS[c][k] for k in cols[3:]] for c, *_ in AGENTS]
    rows += [[p, 4, "PM002", 60_000_000, n, 85, "Active"] for p, n in sorted(PRIOR["PM002"].items())]
    ws = sheet("FAPM_Input", cols, rows)
    for row in ws.iter_rows(min_row=2, max_col=1):
        row[0].number_format = "@"
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=4):
        row[0].number_format = "#,##0"

    keys = ["WorkMonthNo", "PromoMonthNo", "WorkDaysBase", "AttendancePct", "DayFactor", "AllowancePerDay", "K2Coef",
            "CumLPTurns", "LPCountCoef", "FA_Formula", "JoinGuar", "FA_Total", "Reason"]
    ws = sheet("KetQua_KyVong", ["AgentCode", "KichBan"] + keys,
               [[c, note] + [res[c][k] for k in keys] for c, _, _, _, _, note in AGENTS], {1: 48, 14: 34})
    for row in ws.iter_rows(min_row=2, min_col=8, max_col=14):
        for c in row:
            c.number_format = "#,##0.##"
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top")
    wb.save(path)


if __name__ == "__main__":
    main()
