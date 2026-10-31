"""Chay thu FA PM/PD tu file Tbl_Agent + Tbl_Attendance (+ FAPM_Input neu co), cung logic voi modFAPM.bas.

Chay: python3 chay_thu_pm.py <Tbl_Agent.xlsx> <YYYYMM> <ky> <ket_qua.xlsx> [Tbl_Attendance.xlsx] [FAPM_Input.xlsx]
"""

import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "access-fa" / "kiem-tra"))
from chay_thu import load_agents, load_attendance, norm, read_sheet, to_date, to_num  # noqa: E402
from fapm_reference import FA_TITLES, calc  # noqa: E402


def load_inputs(path, period, ky):
    rows, h = read_sheet(load_workbook(path, data_only=True), "FAPM_Input")
    cur, prior = {}, {}
    for r in rows:
        g = lambda k: r[h[k]] if k in h else None
        code, per = str(g("AGENTCODE") or "").strip(), str(g("PERIOD") or "").strip()
        if not code or not per:
            continue
        rec = dict(GroupFYC=to_num(g("GROUPFYC")) or 0, ActiveLP=to_num(g("ACTIVELP")) or 0,
                   K2Pct=to_num(g("K2PCT")), StatusEOM=str(g("STATUSEOM") or "").strip(), Ky=int(g("KY") or 0))
        if per == period and rec["Ky"] == ky:
            cur[code] = rec
        elif per < period:
            # ky lon nhat cua moi thang
            if code not in prior or per not in prior[code] or prior[code][per][1] < rec["Ky"]:
                prior.setdefault(code, {})[per] = (rec["ActiveLP"], rec["Ky"])
    return cur, {c: {p: v for p, (v, _) in d.items()} for c, d in prior.items()}


def main(agent_path, period, ky, out, att_path=None, input_path=None):
    wb = load_workbook(agent_path, data_only=True)
    agents = load_agents(wb, period)
    for code, a in agents.items():
        r = next((x for x in wb["Agent"].iter_rows(min_row=2, values_only=True)
                  if str(x[[norm(h) for h in next(wb["Agent"].iter_rows(max_row=1, values_only=True))].index("AGENTCODE")] or "").strip() == code), None)
        hdr = [norm(h) for h in next(wb["Agent"].iter_rows(max_row=1, values_only=True))]
        a["PromoDate"] = to_date(r[hdr.index("PROMOTIONDATE")]) if r and "PROMOTIONDATE" in hdr else None
    pm = {c: a for c, a in agents.items() if a["Title"] in FA_TITLES}
    print(f"== Agent: {len(agents)} | PM/PD: {len(pm)} ==")
    hol_rows, hh = read_sheet(wb, "Holiday")
    holidays = {to_date(r[hh["HOLIDAYDATE"]]) for r in hol_rows if "HOLIDAYDATE" in hh and to_date(r[hh["HOLIDAYDATE"]])}
    att, dup = load_attendance(att_path, period) if att_path else ({}, 0)
    print(f"== Cham cong {period}: {len(att)} agent, bo {dup} dong trung ==")
    cur, prior = load_inputs(input_path, period, ky) if input_path else ({}, {})
    print(f"== FAPM_Input {period}/{ky}: {len(cur)} dong ==")

    res = []
    for code, a in sorted(pm.items()):
        inp = cur.get(code, dict(GroupFYC=0, ActiveLP=0, K2Pct=None, StatusEOM=a["Status"]))
        inp["ActualDays"] = att.get(code, 0)
        r = calc(period, a, inp, holidays, prior.get(code), ky=ky)
        res.append((code, a["Name"], a["Title"], r))
        print(f"  {code} {a['Name']:<22} {a['Title']:<3} cong={inp['ActualDays']:>4} heso={r['DayFactor']:>5} "
              f"CT={r['FA_Formula']:>11,.0f} baodam6th={r['JoinGuar']:>11,.0f} FA={r['FA_Total']:>12,.0f} {r['Reason']}")
    print(f"  TONG: {sum(r['FA_Total'] for *_, r in res):,.0f}")

    wb2 = Workbook()
    ws = wb2.active
    ws.title = f"FAPM_{period}_Ky{ky}"
    cols = [("Mã", None), ("Họ tên", None), ("Chức danh", None), ("Tháng kể từ thăng tiến", "PromoMonthNo"),
            ("Số ngày LV", "WorkDaysBase"), ("Tỷ lệ chấm công %", "AttendancePct"), ("Hệ số ngày LV", "DayFactor"),
            ("FYC nhóm", "GroupFYC"), ("LP hoạt động", "ActiveLP"), ("Phụ cấp/ngày", "AllowancePerDay"), ("Hệ số K2", "K2Coef"),
            ("Lượt LP lũy kế", "CumLPTurns"), ("Hệ số lượt LP", "LPCountCoef"), ("FA công thức", "FA_Formula"),
            ("Bảo đảm 6 tháng đầu", "JoinGuar"), ("Chi kỳ này", "FA_Total"), ("Ghi chú", "Reason")]
    ws.append([c for c, _ in cols])
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E78")
    for code, name, title, r in res:
        ws.append([code, name, title] + [r.get(k) for _, k in cols[3:]])
    ws.append([None] * 15 + [sum(r["FA_Total"] for *_, r in res)])
    for row in ws.iter_rows(min_row=2, min_col=8, max_col=16):
        for c in row:
            c.number_format = "#,##0.##"
    ws.freeze_panes = "D2"
    wb2.save(out)
    print("->", out)


if __name__ == "__main__":
    a = sys.argv
    main(a[1], a[2], int(a[3]), a[4], a[5] if len(a) > 5 else None, a[6] if len(a) > 6 else None)
