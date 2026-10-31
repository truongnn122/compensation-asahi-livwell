"""Ban tinh doi chieu FA (Muc A.II - ALL/AD/2210-005).

Cung logic voi modFA.bas, dung de:
  1. Kiem tra cong thuc voi vi du minh hoa trong van ban (LP A = 10.500.000d).
  2. Sinh file Excel mau (FA_Input_Mau.xlsx) kem ket qua ky vong de doi soat
     sau khi chay tool Access.

Chay: python3 fa_reference.py
"""

import calendar
from datetime import date, timedelta
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ---- Tham so (khop tblParam trong Build_FA_Tool.vbs) ----
STD_DAYS = 22
MIN_ATT_PCT = 70.0
MIN_DAYS = 15
NEW_GUAR_ALLOW = 300_000
NEW_GUAR_MONTHS = 2
NEW_GUAR_MIN_CT = 0.5
QTR_GUAR_ALLOW = 150_000
QTR_MIN_CT = 6
QTR_MIN_FYC = 24_000_000
K2_NA_COEF = 1.0
FA_TITLES = {"LP"}  # Muc A.II: chi LP toan thoi gian
MDRT_COEF = 1.2

# (Lua chon, so LP gioi thieu toi thieu, FYC toi thieu, phu cap/ngay)
ALLOWANCE = [
    (1, 0, 0, 200_000),
    (1, 0, 15_000_000, 300_000),
    (1, 0, 25_000_000, 500_000),
    (2, 1, 0, 300_000),
    (2, 1, 15_000_000, 500_000),
    (2, 1, 25_000_000, 700_000),
]
K2_TIERS = [(80, 1.2), (70, 1.1), (60, 1.0), (0, 0.5)]
CONTRACT_TIERS = [(2, 1.0), (1, 0.7), (0.5, 0.5), (0, 0.0)]


def rhu(x, d=0):
    """Lam tron half-up (muc D.s), tranh banker's rounding."""
    q = Decimal(1).scaleb(-d)
    return float(Decimal(str(x)).quantize(q, rounding=ROUND_HALF_UP))


def work_days(d1, d2, holidays):
    n, d = 0, d1
    while d <= d2:
        if d.weekday() < 5 and d not in holidays:
            n += 1
        d += timedelta(days=1)
    return n


def allowance(fyc, recruits):
    return max(a for _, r, f, a in ALLOWANCE if fyc >= f and recruits >= r)


def k2_coef(k2):
    if k2 is None:
        return K2_NA_COEF
    k = rhu(k2, 1)
    return next(c for m, c in K2_TIERS if k >= m)


def contract_coef(ct):
    return next(c for m, c in CONTRACT_TIERS if ct >= m)


def calc(period, agent, inp, holidays, m1=None, qtr_qualified=False):
    """Tinh FA thang cho 1 LP. m1 = ket qua thang lam viec thu nhat (neu co)."""
    y, m = int(period[:4]), int(period[4:])
    start = date(y, m, 1)
    end = date(y, m, calendar.monthrange(y, m)[1])
    appoint = agent["AppointDate"]
    wm = (y - appoint.year) * 12 + (m - appoint.month) + 1

    base = work_days(appoint if wm == 1 else start, end, holidays)
    att = min(1.0, (inp["ActualDays"] + 1) / base) if base else 0
    att_pct = rhu(att * 100, 1)
    if wm == 1:
        day_f = rhu(min(inp["ActualDays"], STD_DAYS), 1)
    else:
        day_f = rhu(att_pct / 100 * STD_DAYS, 1)

    reason = ""
    title = (agent.get("Title") or "").strip().upper()
    if title and title not in FA_TITLES:
        reason = "Chức danh không phải LP (FA chỉ áp dụng LP toàn thời gian)"
    elif agent["WorkModel"] != "FT":
        reason = "Không phải LP toàn thời gian"
    elif (inp["StatusEOM"] or "").upper() == "SE":
        reason = "Tình trạng SE cuối tháng"
    elif not inp["TrainingDone"]:
        reason = "Chưa hoàn tất huấn luyện bắt buộc"
    elif att_pct < MIN_ATT_PCT and inp["ActualDays"] < MIN_DAYS:
        reason = "Chấm công < 70% và < 15 ngày"
    eligible = reason == ""

    allow = allowance(inp["FYC"], inp["Recruits"])
    k2c = k2_coef(inp["K2Pct"])
    mdrtc = MDRT_COEF if inp["MDRTPrev"] else 1.0
    ctc = contract_coef(inp["Contracts"])
    fa_formula = rhu(allow * day_f * k2c * mdrtc * ctc)

    new_guar = pending = retro = 0
    if 1 <= wm <= NEW_GUAR_MONTHS:
        cum = inp["Contracts"] + (m1["Contracts"] if wm == 2 and m1 else 0)
        amt = rhu(NEW_GUAR_ALLOW * day_f)
        if eligible:
            if cum >= NEW_GUAR_MIN_CT:
                new_guar = amt
            else:
                pending = amt
        if wm == 2 and m1 and cum >= NEW_GUAR_MIN_CT and m1["NewGuar"] == 0:
            retro = max(0, m1["Pending"] - m1["FAMonth"])

    fa_month = max(fa_formula, new_guar) if eligible else 0
    qtr_month = rhu(QTR_GUAR_ALLOW * day_f) if qtr_qualified and eligible else 0

    return {
        "WorkMonthNo": wm,
        "WorkDaysBase": base,
        "AttendancePct": att_pct,
        "DayFactor": day_f,
        "Eligible": eligible,
        "Reason": reason,
        "AllowancePerDay": allow,
        "K2Coef": k2c,
        "MDRTCoef": mdrtc,
        "ContractCoef": ctc,
        "FA_Formula": fa_formula,
        "NewGuar": new_guar,
        "NewGuarPending": pending,
        "NewGuarRetro": retro,
        "FA_Month": fa_month,
        "QtrGuarMonth": qtr_month,
        "FA_Total": fa_month + retro,
    }


# ---- Du lieu mau: thang 10/2026 (22 ngay lam viec, khong co ngay le) ----
PERIOD = "202610"
HOLIDAYS = {date(2026, 9, 2): "Quốc khánh", date(2027, 1, 1): "Tết Dương lịch"}

AGENTS = [
    # code, ten, ngay bo nhiem, mo hinh, ghi chu kich ban
    ("LP001", "LP A (ví dụ văn bản)", date(2026, 3, 2), "FT", "Ví dụ trang 3: kỳ vọng 10.500.000"),
    ("LP002", "Nguyễn Văn B", date(2025, 6, 2), "FT", "Mức 0,7tr + K2 120% + MDRT 120%"),
    ("LP003", "Trần Thị C", date(2026, 1, 5), "FT", "Mức 0,2tr, 0,5 HĐ → hệ số HĐ 50%"),
    ("LP004", "Lê Văn D", date(2025, 9, 1), "FT", "SE cuối tháng → không chi"),
    ("LP005", "Phạm Thị E", date(2026, 2, 2), "FT", "Chấm công 59,1% và 12 ngày → không chi"),
    ("LP006", "Hoàng Văn F", date(2026, 10, 12), "FT", "LP mới tháng 1, 1 HĐ → bảo đảm 0,3tr"),
    ("LP007", "Vũ Thị G", date(2026, 10, 12), "FT", "LP mới tháng 1, chưa có HĐ → chờ bảo đảm"),
    ("LP008", "Đặng Văn H", date(2025, 11, 3), "PT", "LP bán thời gian → không có FA"),
    ("LP009", "Bùi Thị I", date(2026, 4, 1), "FT", "Chưa hoàn tất huấn luyện → không chi"),
    ("LP010", "Đỗ Văn K", date(2025, 8, 1), "FT", "K2 55% → hệ số K2 50%"),
    ("PM001", "Ngô Thị M", date(2026, 3, 2), "FT", "Chức danh PM → không có FA"),
]
TITLES = {"PM001": "PM"}

# Mau chuan sheet Agent = file xuat he thong Tbl_Agent.xlsx (98 cot, giu nguyen ten va thu tu, ke ca cot trung ten)
AGENT_HEADERS = ["Period", "Channel", "REGION_CODE", "REGION_NAME", "ZONE_CODE", "ZONE_NAME", "BRANCH_CODE", "BRANCH_NAME", "SH_CODE", "SH_NAME", "SD_CODE", "SD_NAME", "INDIRECT 2_NAME", "INDIRECT 2_NAME", "INDIRECT 1_CODE", "INDIRECT 1_NAME", "DIRECT_MANAGER_CODE", "DIRECT_MANAGER_NAME", "LEADER", "EXT_REF", "AGENT_CODE", "FULL_NAME", "GENDER", "TITLE (Code)", "AGENT_STATUS", "DOB", "NATIONAL_ID", "ID_ISSUE_DATE", "ID_ISSUED_BY", "ID_EXPIRATION_DATE", "OLD_ID_NO", "PLACE_OF_BIRTH", "NATIONALITY", "MARITAL_STATUS", "TAX_CODE", "INCOME_RANGE", "EDUCATION", "PUBLIC_SECTOR", "OCCUPATION", "EXPERIENCE (Year)", "PERM_PROVINCE", "PERM_WARD", "PERM_STREET", "CONTACT_PROVINCE", "CONTACT_WARD", "CONTACT_STREET", "MOBILE", "PERSONAL_EMAIL", "COMPANY_EMAIL", "CHANNEL", "AGENCY_TYPE", "FIRST_RECRUITER_CODE", "REFERRER_CODE", "APPOINTMENT_DATE", "TERMINATION_DATE", "FIRST APPOINTMENT_DATE", "PROMOTION_DATE", "DEMOTION_DATE", "BANK_ACCOUNT_HOLDER", "BANK_ACCOUNT_NO", "BANK_CODE", "BANK_BRANCH", "BANK_BRANCH CODE (Payment)", "MOF_CERT_NO", "MOF_CERT_ISSUE_DATE", "MOF_EXAM_PASS_DATE", "AVICAD_STATUS", "AVICAD_DESC", "AVICAD_CHECKED_AT", "CONTRACT_NO Agent", "CONTRACT_TYPE", "CONTRACT_NO Leader", "CONTRACT_TYPE", "ESIGN_ENVELOPE_ID (DOCUMENT ID)", "CONTRACT_SENT_DATE", "CONTRACT_SIGNED_DATE", "SIGN_AUTH_METHOD", "SIGN_STATUS", "CONTRACT_FILE_NAME", "CONTRACT_DMS_REF", "PROGRAM_CODES", "COMMIT_1", "COMMIT_2", "COMMIT_3", "COMMIT_4", "CONSENT_NOTIFY_OPS", "CONSENT_NOTIFY_PROMO", "PDPD_1_BASIC", "PDPD_2_SENSITIVE", "PDPD_3_SHARE", "CONSENT_DATE", "CONSENT_METHOD", "ONBOARD_COURSE_CODE", "ONBOARD_COURSE_DATE", "TRAINING_COMPLETED", "PRODUCT_AUTH_CODES", "DOC_RECEIVED", "DOC_DMS_REF"]
DATE_HEADERS = {h for h in AGENT_HEADERS if "DATE" in h.upper() or h.upper() == "DOB" or h.upper().endswith("_AT")}

INPUTS = {
    "LP001": dict(ActualDays=20, FYC=26_000_000, Contracts=2, Recruits=0, K2Pct=None, MDRTPrev=0, TrainingDone=1, StatusEOM="Active"),
    "LP002": dict(ActualDays=22, FYC=30_000_000, Contracts=3, Recruits=1, K2Pct=85, MDRTPrev=1, TrainingDone=1, StatusEOM="Active"),
    "LP003": dict(ActualDays=16, FYC=8_000_000, Contracts=0.5, Recruits=0, K2Pct=65, MDRTPrev=0, TrainingDone=1, StatusEOM="Active"),
    "LP004": dict(ActualDays=18, FYC=5_000_000, Contracts=1, Recruits=0, K2Pct=72, MDRTPrev=0, TrainingDone=1, StatusEOM="SE"),
    "LP005": dict(ActualDays=12, FYC=16_000_000, Contracts=2, Recruits=0, K2Pct=None, MDRTPrev=0, TrainingDone=1, StatusEOM="Active"),
    "LP006": dict(ActualDays=14, FYC=6_000_000, Contracts=1, Recruits=0, K2Pct=None, MDRTPrev=0, TrainingDone=1, StatusEOM="Active"),
    "LP007": dict(ActualDays=13, FYC=0, Contracts=0, Recruits=0, K2Pct=None, MDRTPrev=0, TrainingDone=1, StatusEOM="Active"),
    "LP008": dict(ActualDays=10, FYC=12_000_000, Contracts=1, Recruits=0, K2Pct=75, MDRTPrev=0, TrainingDone=1, StatusEOM="Active"),
    "LP009": dict(ActualDays=21, FYC=20_000_000, Contracts=2, Recruits=0, K2Pct=None, MDRTPrev=0, TrainingDone=0, StatusEOM="Active"),
    "PM001": dict(ActualDays=20, FYC=26_000_000, Contracts=2, Recruits=0, K2Pct=None, MDRTPrev=0, TrainingDone=1, StatusEOM="Active"),
    "LP010": dict(ActualDays=19, FYC=18_000_000, Contracts=2, Recruits=1, K2Pct=55, MDRTPrev=0, TrainingDone=1, StatusEOM="Active"),
}


def main():
    hol = set(HOLIDAYS)
    results = {}
    for code, _, appoint, model, _ in AGENTS:
        agent = {"AppointDate": appoint, "WorkModel": model, "Title": TITLES.get(code, "LP")}
        results[code] = calc(PERIOD, agent, INPUTS[code], hol)

    assert results["LP001"]["FA_Total"] == 10_500_000, results["LP001"]
    assert results["LP002"]["FA_Total"] == 22_176_000, results["LP002"]
    assert results["LP003"]["FA_Total"] == 1_700_000, results["LP003"]
    assert results["LP006"]["FA_Total"] == 4_200_000, results["LP006"]
    assert results["LP007"]["NewGuarPending"] == 3_900_000, results["LP007"]
    assert results["LP010"]["FA_Total"] == 5_000_000, results["LP010"]
    for c in ("LP004", "LP005", "LP008", "LP009", "LP007", "PM001"):
        assert results[c]["FA_Total"] == 0, (c, results[c])

    out = Path(__file__).resolve().parent.parent / "FA_Input_Mau.xlsx"
    build_workbook(out, results)
    for code, *_ in AGENTS:
        r = results[code]
        print(f"{code}: FA={r['FA_Total']:>12,.0f}  {r['Reason']}")
    print("OK ->", out)


def build_workbook(path, results):
    wb = Workbook()
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1F4E78")

    def sheet(title, cols, rows, widths=None):
        ws = wb.create_sheet(title)
        ws.append(cols)
        for c in ws[1]:
            c.font, c.fill = head, fill
        for r in rows:
            ws.append(list(r))
        for i, _ in enumerate(cols):
            ws.column_dimensions[get_column_letter(i + 1)].width = (widths or {}).get(i, 16)
        ws.freeze_panes = "A2"
        return ws

    hd = wb.active
    hd.title = "HuongDan"
    guide = [
        "MẪU NHẬP DỮ LIỆU TÍNH PHỤ CẤP CỐ ĐỊNH (FA) – ALL/AD/2210-005 mục A.II",
        "",
        "Tool chỉ đọc 3 sheet: Agent, Holiday, FA_Input. Không đổi tên sheet và tên cột.",
        "",
        "Sheet Agent – danh sách agent hàng tháng: dán NGUYÊN file xuất hệ thống (Tbl_Agent, 98 cột, giữ đúng tên và thứ tự cột)",
        "  Cột bắt buộc: Period, AGENT_CODE, FULL_NAME, TITLE (Code), AGENCY_TYPE (FT/PT), APPOINTMENT_DATE",
        "  TERMINATION_DATE để trống nếu HĐ còn hiệu lực. FA chỉ tính cho chức danh LP",
        "  Nạp lại cùng Period thì ghi đè đúng tháng đó; các tháng khác giữ nguyên trong lịch sử",
        "",
        "Sheet Holiday – ngày lễ rơi vào ngày làm việc (để tính số ngày làm việc của tháng)",
        "",
        "Sheet FA_Input – số liệu từng LP theo tháng và kỳ chi trả",
        "  Period: tháng tính, dạng YYYYMM (ví dụ 202610)",
        "  Ky: kỳ chi trả 1–4 (mục D.l). Mỗi kỳ nhập số LŨY KẾ đến ngày chốt ACK của kỳ đó",
        "  ActualDays: số ngày đi làm thực tế trong tháng (gồm cả ngày làm thêm theo yêu cầu)",
        "  FYC: FYC phát sinh (VND) | Contracts: số hợp đồng phát sinh (cho phép 0,5)",
        "  Recruits: số LP giới thiệu trong tháng",
        "  K2Pct: %K2 dạng số, ví dụ 85 = 85%. Để TRỐNG nếu chưa có K2 (N/A)",
        "  MDRTPrev: 1 = đạt MDRT năm liền kề trước, 0 = không",
        "  TrainingDone: 1 = đã hoàn tất huấn luyện bắt buộc, 0 = chưa",
        "  StatusEOM: tình trạng cuối tháng (Active / SE / ...)",
        "",
        "Sheet KetQua_KyVong: kết quả tính tay bằng bản đối chiếu Python – dùng để đối soát với tool, tool không đọc sheet này.",
    ]
    for line in guide:
        hd.append([line])
    hd["A1"].font = Font(bold=True, size=13)
    hd.column_dimensions["A"].width = 120

    def agent_row(code, name, appoint, model):
        v = {"Period": PERIOD, "Channel": "AGENCY", "AGENT_CODE": code, "FULL_NAME": name,
             "TITLE (Code)": TITLES.get(code, "LP"), "AGENT_STATUS": "Active", "CHANNEL": "AGENCY",
             "AGENCY_TYPE": model, "APPOINTMENT_DATE": appoint}
        return [v.get(h) for h in AGENT_HEADERS]

    ws = sheet("Agent", AGENT_HEADERS, [agent_row(c, n, d, m) for c, n, d, m, _ in AGENTS])
    for i, h in enumerate(AGENT_HEADERS):
        if h in DATE_HEADERS:
            for row in ws.iter_rows(min_row=2, min_col=i + 1, max_col=i + 1):
                row[0].number_format = "dd/mm/yyyy"
        if h in ("Period", "AGENT_CODE"):
            for row in ws.iter_rows(min_row=2, min_col=i + 1, max_col=i + 1):
                row[0].number_format = "@"

    ws = sheet("Holiday", ["HolidayDate", "Description"], HOLIDAYS.items(), {1: 30})
    for row in ws.iter_rows(min_row=2, max_col=1):
        row[0].number_format = "dd/mm/yyyy"

    cols = ["Period", "Ky", "AgentCode", "ActualDays", "FYC", "Contracts",
            "Recruits", "K2Pct", "MDRTPrev", "TrainingDone", "StatusEOM"]
    ws = sheet(
        "FA_Input",
        cols,
        [[PERIOD, 1, c] + [INPUTS[c][k] for k in cols[3:]] for c, *_ in AGENTS],
    )
    for row in ws.iter_rows(min_row=2, max_col=1):
        row[0].number_format = "@"
    for row in ws.iter_rows(min_row=2, min_col=5, max_col=5):
        row[0].number_format = "#,##0"

    keys = ["WorkMonthNo", "WorkDaysBase", "AttendancePct", "DayFactor",
            "AllowancePerDay", "K2Coef", "MDRTCoef", "ContractCoef",
            "FA_Formula", "NewGuar", "NewGuarPending", "FA_Total", "Reason"]
    ws = sheet(
        "KetQua_KyVong",
        ["AgentCode", "KichBan"] + keys,
        [[c, note] + [results[c][k] for k in keys] for c, _, _, _, note in AGENTS],
        {1: 40, 14: 34},
    )
    for row in ws.iter_rows(min_row=2, min_col=7, max_col=13):
        for c in row:
            c.number_format = "#,##0.##"
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top")
    wb.save(path)


if __name__ == "__main__":
    main()
