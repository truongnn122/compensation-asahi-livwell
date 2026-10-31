"""Chay thu tinh FA tu file Excel nguoi dung, cung logic voi modFA.bas.

Doc sheet Agent (mau tool hoac file xuat he thong), Holiday, FA_Input; tinh FA cho
thang/ky chi dinh; in tung buoc kiem tra va xuat bang ke Excel.

Chay: python3 chay_thu.py <file.xlsx> <YYYYMM> <ky> <file_ket_qua.xlsx>
"""

import sys
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from fa_reference import calc


def norm(h):
    return str(h or "").strip().replace(" ", "").replace("_", "").upper()


def read_sheet(wb, name):
    if name not in wb.sheetnames:
        return [], {}
    rows = [r for r in wb[name].iter_rows(values_only=True)]
    if not rows:
        return [], {}
    hdr = {}
    for i, h in enumerate(rows[0]):
        k = norm(h)
        if k and k not in hdr:
            hdr[k] = i
    return rows[1:], hdr


def col(hdr, *names):
    return next((hdr[n] for n in names if n in hdr), None)


def to_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return None


def to_num(v):
    if v is None or str(v).strip() == "":
        return None
    return float(v)


def load_agents(wb):
    rows, hdr = read_sheet(wb, "Agent")
    c = dict(
        code=col(hdr, "AGENTCODE"), name=col(hdr, "AGENTNAME", "FULLNAME"),
        appoint=col(hdr, "APPOINTDATE", "APPOINTMENTDATE"), model=col(hdr, "WORKMODEL", "AGENCYTYPE"),
        term=col(hdr, "TERMINATEDATE", "TERMINATIONDATE"), title=col(hdr, "TITLE", "TITLE(CODE)"),
        status=col(hdr, "AGENTSTATUS"),
    )
    missing = [k for k in ("code", "appoint", "model") if c[k] is None]
    if missing:
        sys.exit(f"Sheet Agent thieu cot bat buoc: {missing}")
    agents = {}
    for r in rows:
        code = str(r[c["code"]] or "").strip()
        if not code:
            continue
        get = lambda k: r[c[k]] if c[k] is not None else None
        agents[code] = dict(
            Name=str(get("name") or "").strip(), AppointDate=to_date(get("appoint")),
            WorkModel=str(get("model") or "").strip().upper()[:2], TerminateDate=to_date(get("term")),
            Title=str(get("title") or "").strip().upper(), Status=str(get("status") or "").strip(),
        )
    return agents


def main(path, period, ky, out):
    wb = load_workbook(path, data_only=True)

    print("== BUOC 1: Danh sach agent ==")
    agents = load_agents(wb)
    by_title = {}
    for a in agents.values():
        by_title[a["Title"] or "(trong)"] = by_title.get(a["Title"] or "(trong)", 0) + 1
    print(f"  {len(agents)} agent | theo chuc danh: {by_title}")
    for code, a in agents.items():
        if not a["AppointDate"]:
            print(f"  ! {code}: thieu ngay bo nhiem")

    hol_rows, hh = read_sheet(wb, "Holiday")
    holidays = {to_date(r[hh["HOLIDAYDATE"]]) for r in hol_rows if "HOLIDAYDATE" in hh and to_date(r[hh["HOLIDAYDATE"]])}
    print(f"  {len(holidays)} ngay le")

    print(f"== BUOC 2: Du lieu FA_Input thang {period} ky {ky} ==")
    rows, h = read_sheet(wb, "FA_Input")
    inputs = {}
    for r in rows:
        g = lambda k: r[h[k]] if k in h else None
        if str(g("PERIOD") or "").strip() != period or int(g("KY") or 0) != ky:
            continue
        code = str(g("AGENTCODE") or "").strip()
        if code:
            inputs[code] = dict(
                ActualDays=to_num(g("ACTUALDAYS")) or 0, FYC=to_num(g("FYC")) or 0,
                Contracts=to_num(g("CONTRACTS")) or 0, Recruits=to_num(g("RECRUITS")) or 0,
                K2Pct=to_num(g("K2PCT")), MDRTPrev=(to_num(g("MDRTPREV")) or 0) == 1,
                TrainingDone=(to_num(g("TRAININGDONE")) or 0) == 1, StatusEOM=str(g("STATUSEOM") or "").strip(),
            )
    matched = [c for c in inputs if c in agents]
    print(f"  {len(inputs)} dong | khop danh sach agent: {len(matched)}")
    for c in inputs:
        if c not in agents:
            print(f"  ! {c}: khong co trong sheet Agent")
    for code, a in agents.items():
        if a["Title"] in ("", "LP") and code not in inputs:
            print(f"  ! {code} ({a['Name']}, {a['Title'] or '?'}): chua co dong FA_Input")

    print("== BUOC 3: Tinh FA ==")
    res = []
    for code, inp in inputs.items():
        a = agents.get(code)
        if not a or not a["AppointDate"]:
            res.append((code, a["Name"] if a else "", "", None, "Không có trong danh sách Agent / thiếu ngày bổ nhiệm"))
            continue
        r = calc(period, a, inp, holidays)
        if a["TerminateDate"] and r["Eligible"] and a["TerminateDate"].strftime("%Y%m") <= period:
            r.update(Eligible=False, FA_Month=0, FA_Total=0, Reason="HĐ đại lý chấm dứt trong/trước tháng tính")
        if r["NewGuarPending"] and not r["Reason"]:
            r["Reason"] = "LP mới chưa đủ 0,5 HĐ – bảo đảm chờ"
        res.append((code, a["Name"], a["Title"], r, r["Reason"]))
        print(f"  {code} {a['Name']:<22} {a['Title']:<3} FA={r['FA_Total']:>12,.0f}  {r['Reason']}")
    total = sum(r[3]["FA_Total"] for r in res if r[3])
    print(f"  TONG: {total:,.0f}")
    write_result(out, period, ky, res)
    print("-> ", out)


def write_result(path, period, ky, res):
    wb = Workbook()
    ws = wb.active
    ws.title = f"FA_{period}_Ky{ky}"
    cols = [("Mã LP", None), ("Họ tên", None), ("Chức danh", None),
            ("Tháng LV thứ", "WorkMonthNo"), ("Số ngày LV của tháng", "WorkDaysBase"),
            ("Tỷ lệ chấm công %", "AttendancePct"), ("Hệ số ngày LV", "DayFactor"),
            ("Phụ cấp/ngày", "AllowancePerDay"), ("Hệ số K2", "K2Coef"), ("Hệ số MDRT", "MDRTCoef"),
            ("Hệ số HĐ", "ContractCoef"), ("FA công thức", "FA_Formula"), ("Bảo đảm LP mới", "NewGuar"),
            ("Bảo đảm chờ HĐ", "NewGuarPending"), ("Chi kỳ này", "FA_Total"), ("Ghi chú", None)]
    ws.append([c for c, _ in cols])
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E78")
    for code, name, title, r, note in res:
        ws.append([code, name, title] + [(r or {}).get(k) for _, k in cols[3:-1]] + [note])
    ws.append([None] * 13 + ["Tổng", sum((r or {}).get("FA_Total", 0) for *_, r, _ in res)])
    for i, w in enumerate([11, 24, 9, 9, 10, 10, 9, 12, 8, 8, 8, 13, 13, 13, 13, 40]):
        ws.column_dimensions[chr(65 + i)].width = w
    for row in ws.iter_rows(min_row=2, min_col=8, max_col=15):
        for c in row:
            c.number_format = "#,##0.##"
    ws.freeze_panes = "D2"
    wb.save(path)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2], int(sys.argv[3]), sys.argv[4])
